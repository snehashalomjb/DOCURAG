import os
import pickle
import logging
from typing import List, Dict, Any, Tuple
import numpy as np
import faiss
from sentence_transformers import SentenceTransformer
import google.generativeai as genai
from pypdf import PdfReader

from app.backend import config

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Lazy loaded embedding model
_embedding_model = None

def get_embedding_model() -> SentenceTransformer:
    """Lazy load the embedding model to optimize startup time and memory."""
    global _embedding_model
    if _embedding_model is None:
        logger.info(f"Loading sentence-transformer model: {config.EMBEDDING_MODEL_NAME}...")
        _embedding_model = SentenceTransformer(config.EMBEDDING_MODEL_NAME)
        logger.info("Embedding model loaded successfully.")
    return _embedding_model

def configure_gemini():
    """Configure the Gemini API using key from config."""
    if not config.GEMINI_API_KEY:
        logger.warning("GEMINI_API_KEY is not set. Gemini calls will fail until configured.")
    genai.configure(api_key=config.GEMINI_API_KEY)

class RAGPipeline:
    def __init__(self):
        self.index = None
        self.chunks = []
        self.chunk_sources = []  # Tracks which file each chunk came from
        self.metadata = {"documents": []}
        configure_gemini()

    def load_document(self, file_path: str) -> str:
        """Extracts text content from a PDF, TXT, DOCX, HTML, or MD file."""
        logger.info(f"Loading document: {file_path}")
        _, ext = os.path.splitext(file_path.lower())
        
        if ext == '.pdf':
            return self._extract_pdf(file_path)
        elif ext == '.txt':
            return self._extract_txt(file_path)
        elif ext == '.docx':
            return self._extract_docx(file_path)
        elif ext in ['.html', '.htm']:
            return self._extract_html(file_path)
        elif ext == '.md':
            return self._extract_md(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: PDF, TXT, DOCX, HTML, MD.")

    def _extract_pdf(self, file_path: str) -> str:
        text = []
        try:
            reader = PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
            full_text = "\n\n".join(text)
            if not full_text.strip():
                raise ValueError("The PDF file appears to contain no text or only scanned images (OCR not supported).")
            return full_text
        except Exception as e:
            logger.error(f"Error reading PDF: {e}")
            raise RuntimeError(f"Failed to read PDF: {str(e)}")

    def _extract_txt(self, file_path: str) -> str:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error reading TXT file: {e}")
            raise RuntimeError(f"Failed to read TXT file: {str(e)}")

    def _extract_docx(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            text = []
            for para in doc.paragraphs:
                if para.text:
                    text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells if cell.text]
                    if row_text:
                        text.append(" | ".join(row_text))
            full_text = "\n".join(text)
            if not full_text.strip():
                raise ValueError("The DOCX file appears to contain no text.")
            return full_text
        except Exception as e:
            logger.error(f"Error reading DOCX: {e}")
            raise RuntimeError(f"Failed to read DOCX: {str(e)}")

    def _extract_html(self, file_path: str) -> str:
        try:
            from bs4 import BeautifulSoup
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
                
            text = soup.get_text(separator='\n')
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            full_text = '\n'.join(chunk for chunk in chunks if chunk)
            
            if not full_text.strip():
                raise ValueError("The HTML file contains no extractable text.")
            return full_text
        except Exception as e:
            logger.error(f"Error reading HTML: {e}")
            raise RuntimeError(f"Failed to read HTML: {str(e)}")

    def _extract_md(self, file_path: str) -> str:
        try:
            return self._extract_txt(file_path)
        except Exception as e:
            logger.error(f"Error reading MD: {e}")
            raise RuntimeError(f"Failed to read MD: {str(e)}")

    def chunk_text(self, text: str, chunk_size: int = None, chunk_overlap: int = None) -> List[str]:
        """Splits text into chunks using overlapping sliding window."""
        if chunk_size is None:
            chunk_size = config.DEFAULT_CHUNK_SIZE
        if chunk_overlap is None:
            chunk_overlap = config.DEFAULT_CHUNK_OVERLAP

        if chunk_size <= 0:
            raise ValueError("chunk_size must be positive.")
        if chunk_overlap >= chunk_size:
            raise ValueError("chunk_overlap must be less than chunk_size.")

        logger.info(f"Chunking text with chunk_size={chunk_size}, chunk_overlap={chunk_overlap}")
        
        # Clean text slightly
        text = " ".join(text.split())
        
        chunks = []
        start = 0
        text_len = len(text)
        
        if text_len == 0:
            return []

        while start < text_len:
            end = start + chunk_size
            
            # If we aren't at the end of the text, try to find a space to break on
            if end < text_len:
                # Search backwards for a space within 20% of the chunk size
                limit = max(start, end - int(chunk_size * 0.2))
                space_idx = text.rfind(' ', limit, end)
                if space_idx != -1:
                    end = space_idx + 1

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
                
            start = end - chunk_overlap
            if start >= text_len or (end >= text_len):
                break
                
        logger.info(f"Created {len(chunks)} chunks.")
        return chunks

    def build_index(self, chunks: List[str]):
        """Generates embeddings and builds a brand new FAISS index (legacy/single-file usage)."""
        if not chunks:
            raise ValueError("No text chunks provided for index building.")
            
        logger.info(f"Generating embeddings for {len(chunks)} chunks...")
        model = get_embedding_model()
        
        embeddings = model.encode(chunks, show_progress_bar=False)
        embeddings_np = np.array(embeddings).astype('float32')
        
        dimension = config.EMBEDDING_DIMENSION
        self.index = faiss.IndexFlatL2(dimension)
        self.index.add(embeddings_np)
        self.chunks = chunks
        self.chunk_sources = ["Unknown"] * len(chunks)
        self.metadata = {"documents": ["Unknown"]}
        
        logger.info("FAISS Index built successfully.")

    def append_to_index(self, chunks: List[str], source_filename: str):
        """Generates embeddings for new chunks and appends them to the active FAISS index."""
        if not chunks:
            raise ValueError("No text chunks provided to append.")
            
        logger.info(f"Generating embeddings for {len(chunks)} new chunks from '{source_filename}'...")
        model = get_embedding_model()
        
        embeddings = model.encode(chunks, show_progress_bar=False)
        embeddings_np = np.array(embeddings).astype('float32')
        
        dimension = config.EMBEDDING_DIMENSION
        
        # Initialize index if it doesn't exist yet
        if self.index is None:
            self.index = faiss.IndexFlatL2(dimension)
            self.chunks = []
            self.chunk_sources = []
            self.metadata = {"documents": []}
            
        self.index.add(embeddings_np)
        self.chunks.extend(chunks)
        self.chunk_sources.extend([source_filename] * len(chunks))
        
        if "documents" not in self.metadata:
            self.metadata["documents"] = []
        if source_filename not in self.metadata["documents"]:
            self.metadata["documents"].append(source_filename)
            
        logger.info(f"FAISS index appended. Chunks total: {len(self.chunks)}.")

    def generate_summary(self, text: str) -> str:
        """Generates a professional, grounded 3-line bulleted summary using Gemini."""
        if not config.GEMINI_API_KEY:
            return "Gemini API Key not set. Could not generate summary."
        if not text.strip():
            return "Empty document. Nothing to summarize."

        # Grab a representative snippet from the start (up to 3000 chars)
        snippet = text[:3000].strip()
        
        prompt = (
            "You are a professional document analysis assistant. Read the following document excerpt and generate "
            "a concise summary in exactly 3 bullet points.\n"
            "Rules:\n"
            "1. Output exactly 3 bullet points.\n"
            "2. Keep each bullet point under 15 words.\n"
            "3. Focus on identifying the primary topic, key theme, and structure of the document.\n"
            "4. Do not output any preamble, greeting, or explanations. Only the 3 bullet points.\n\n"
            f"--- EXCERPT CONTENT ---\n{snippet}\n"
            "--- END EXCERPT CONTENT ---"
        )
        
        try:
            configure_gemini()
            model = genai.GenerativeModel(model_name=config.GEMINI_MODEL_NAME)
            logger.info("Calling Gemini for auto-summary generation...")
            res = model.generate_content(
                prompt,
                generation_config=genai.types.GenerationConfig(temperature=0.2)
            )
            summary = res.text.strip()
            if not summary:
                summary = "- Summary generation returned empty text."
            return summary
        except Exception as e:
            logger.error(f"Failed to generate summary: {e}")
            return f"- Failed to generate summary due to error: {str(e)}"

    def generate_analytics(self, text: str, file_path: str) -> Dict[str, Any]:
        """Generates statistical and semantic insights (pages, words, topics, questions) from a document using Gemini."""
        # Calculate word count
        word_count = len(text.split())
        
        # Calculate page count
        page_count = 1
        _, ext = os.path.splitext(file_path.lower())
        if ext == '.pdf':
            try:
                reader = PdfReader(file_path)
                page_count = len(reader.pages)
            except Exception as e:
                logger.warning(f"Failed to read PDF page count: {e}")
                page_count = max(1, int(len(text) / 1500))
        elif ext == '.docx':
            try:
                from docx import Document
                doc = Document(file_path)
                page_count = max(1, int(word_count / 450))
            except Exception:
                page_count = max(1, int(word_count / 450))
        else:
            page_count = max(1, int(word_count / 500))

        # Default fallback analytics
        analytics = {
            "page_count": page_count,
            "word_count": word_count,
            "key_topics": ["General Information"],
            "suggested_questions": [
                "What is the main topic of this document?", 
                "Can you summarize the key findings?", 
                "What are the most important sections?"
            ]
        }

        if not config.GEMINI_API_KEY:
            return analytics

        snippet = text[:3000].strip()
        prompt = (
            "You are a professional document analyst. Review the following text excerpt of a document and extract:\n"
            "1. The top 5 key topics (short phrases, maximum 4 words each).\n"
            "2. Exactly 3 suggested questions that a reader might want to ask to probe the document's core content.\n\n"
            "You MUST respond in JSON format matching this schema:\n"
            "{\n"
            "  \"key_topics\": [\"Topic 1\", \"Topic 2\", \"Topic 3\", \"Topic 4\", \"Topic 5\"],\n"
            "  \"suggested_questions\": [\"Question 1\", \"Question 2\", \"Question 3\"]\n"
            "}\n\n"
            f"--- EXCERPT CONTENT ---\n{snippet}\n"
            "--- END EXCERPT CONTENT ---"
        )

        try:
            configure_gemini()
            model = genai.GenerativeModel(model_name=config.GEMINI_MODEL_NAME)
            logger.info("Calling Gemini with response_mime_type=application/json for structured analytics...")
            
            res = model.generate_content(
                prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.2,
                    response_mime_type="application/json"
                )
            )
            
            import json
            data = json.loads(res.text.strip())
            
            analytics["key_topics"] = data.get("key_topics", analytics["key_topics"])[:5]
            analytics["suggested_questions"] = data.get("suggested_questions", analytics["suggested_questions"])[:3]
            logger.info("Structured analytics generated successfully.")
        except Exception as e:
            logger.error(f"Failed to generate structured analytics via Gemini: {e}")
            
        return analytics


    def search(self, query: str, top_k: int = None) -> List[Dict[str, Any]]:
        """Retrieves top-k closest chunks from the index, calculating match confidence classes."""
        if self.index is None or not self.chunks:
            raise ValueError("No documents indexed for this session.")
            
        if top_k is None:
            top_k = config.DEFAULT_TOP_K
            
        top_k = min(top_k, len(self.chunks))
        if top_k <= 0:
            return []

        logger.info(f"Searching index for query: '{query}' with top_k={top_k}")
        model = get_embedding_model()
        query_embedding = model.encode([query]).astype('float32')
        
        distances, indices = self.index.search(query_embedding, top_k)
        
        results = []
        for i, idx in enumerate(indices[0]):
            if idx != -1 and idx < len(self.chunks):
                dist = float(distances[0][i])
                
                # Confidence Category Mapping based on L2 distance
                if dist < 0.8:
                    confidence = "High"
                elif dist < 1.3:
                    confidence = "Medium"
                else:
                    confidence = "Low"
                    
                results.append({
                    "chunk": self.chunks[idx],
                    "score": dist,
                    "confidence": confidence,
                    "chunk_index": int(idx),
                    "source_doc": self.chunk_sources[idx] if idx < len(self.chunk_sources) else "Unknown"
                })
        return results

    def save_index(self, name: str):
        """Saves current state (FAISS index + chunks + sources + metadata) to disk."""
        if self.index is None:
            raise ValueError("No index built to save.")
            
        logger.info(f"Saving index '{name}' to disk...")
        index_file = os.path.join(config.INDEX_DIR, f"{name}.faiss")
        meta_file = os.path.join(config.INDEX_DIR, f"{name}.pkl")
        
        faiss.write_index(self.index, index_file)
        
        state = {
            "chunks": self.chunks,
            "chunk_sources": self.chunk_sources,
            "metadata": self.metadata
        }
        with open(meta_file, 'wb') as f:
            pickle.dump(state, f)
        logger.info("Index and metadata saved successfully.")

    def load_index(self, name: str) -> bool:
        """Loads index and metadata from disk. Returns True if successful."""
        index_file = os.path.join(config.INDEX_DIR, f"{name}.faiss")
        meta_file = os.path.join(config.INDEX_DIR, f"{name}.pkl")
        
        if not (os.path.exists(index_file) and os.path.exists(meta_file)):
            logger.info(f"No saved index found at '{index_file}' or '{meta_file}'. Cleaning memory.")
            self.index = None
            self.chunks = []
            self.chunk_sources = []
            self.metadata = {"documents": []}
            return False
            
        try:
            logger.info(f"Loading index '{name}' from disk...")
            self.index = faiss.read_index(index_file)
            
            with open(meta_file, 'rb') as f:
                state = pickle.load(f)
                self.chunks = state["chunks"]
                self.chunk_sources = state.get("chunk_sources", ["Unknown"] * len(self.chunks))
                self.metadata = state.get("metadata", {"documents": ["Unknown"]})
            logger.info("Index loaded successfully.")
            return True
        except Exception as e:
            logger.error(f"Error loading index: {e}")
            return False

    def condense_query(self, query: str, history: List[Dict[str, str]]) -> str:
        """Rewrites a follow-up user query using chat history into a standalone search query."""
        if not history or not config.GEMINI_API_KEY:
            return query

        history_str = ""
        for msg in history[-5:]:  # Limit to last 5 turns to stay fast and within limits
            role = "User" if msg["role"] == "user" else "Assistant"
            history_str += f"{role}: {msg['content']}\n"

        rewrite_prompt = (
            "You are an AI assistant. Given the following conversation history and a follow-up query from the user, "
            "your job is to rewrite the follow-up query into a standalone, self-contained search query. "
            "The standalone query will be used for document retrieval in a vector database.\n"
            "Rules:\n"
            "1. Do not answer the question; only rewrite it.\n"
            "2. Ensure all pronouns (he, she, it, they, that, this, the first point) are resolved to their actual concepts mentioned in the history.\n"
            "3. If the follow-up query is already self-contained, return it exactly as is without any explanation or introduction.\n"
            "4. Keep the output query concise and direct.\n\n"
            f"--- CONVERSATION HISTORY ---\n{history_str}\n"
            f"--- FOLLOW-UP QUERY ---\n{query}\n\n"
            "Standalone Query:"
        )

        try:
            configure_gemini()
            model = genai.GenerativeModel(model_name=config.GEMINI_MODEL_NAME)
            response = model.generate_content(
                rewrite_prompt,
                generation_config=genai.types.GenerationConfig(temperature=0.0)  # Greedy decoding
            )
            rewritten = response.text.strip()
            if rewritten:
                logger.info(f"Condense query: '{query}' -> '{rewritten}'")
                return rewritten
        except Exception as e:
            logger.error(f"Error condensing query: {e}")
        
        return query

    def generate_answer(self, query: str, history: List[Dict[str, str]], retrieved_chunks: List[Dict[str, Any]]) -> Tuple[str, List[str]]:
        """Sends query, conversation history, and grounded context to Gemini API to return a response."""
        if not config.GEMINI_API_KEY:
            return "Error: Gemini API Key is not set. Please set the GEMINI_API_KEY environment variable or provide it in the UI settings.", []

        if not retrieved_chunks:
            return "No relevant context found in the loaded document to answer your question.", []

        # Build context
        context_str = ""
        source_texts = []
        for i, res in enumerate(retrieved_chunks):
            chunk_text = res["chunk"]
            source_doc = res.get("source_doc", "Unknown")
            source_texts.append(f"({source_doc}) {chunk_text}")
            context_str += f"[Context Block {i+1} (Source File: {source_doc})]:\n{chunk_text}\n\n"

        # Format history for conversational context
        history_str = ""
        for msg in history[-6:]:  # Keep up to 6 turns for reasoning context
            role = "User" if msg["role"] == "user" else "Assistant"
            history_str += f"{role}: {msg['content']}\n"

        system_prompt = (
            "You are a helpful, professional, and grounded AI Document Assistant. "
            "Your task is to answer the user's question using ONLY the provided document context below.\n"
            "Strict rules:\n"
            "1. Answer the question accurately and comprehensively using ONLY the facts explicitly mentioned in the context.\n"
            "2. If the answer cannot be fully derived from the provided context, state clearly: "
            "'I cannot answer this question based on the uploaded document, as the context does not contain this information.' "
            "Do not try to make up information or use external pre-trained knowledge for facts that aren't in the context.\n"
            "3. Be direct, clear, and refer back to facts/sections in the context when appropriate.\n"
            "4. Take the recent conversation history into account if the user refers to it, but maintain all answers strictly grounded on the context.\n\n"
            f"--- DOCUMENT CONTEXT ---\n{context_str}\n"
            "--- END OF DOCUMENT CONTEXT ---"
        )

        user_input_content = f"{history_str}User: {query}" if history_str else query

        try:
            configure_gemini()
            model = genai.GenerativeModel(
                model_name=config.GEMINI_MODEL_NAME,
                system_instruction=system_prompt
            )
            
            logger.info(f"Calling Gemini ({config.GEMINI_MODEL_NAME}) with user query...")
            response = model.generate_content(
                user_input_content,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.2,  # Low temperature for factual grounding
                )
            )
            return response.text, source_texts
        except Exception as e:
            logger.error(f"Error generating content from Gemini: {e}")
            return f"Error communicating with Gemini API: {str(e)}", []
